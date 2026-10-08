"""Processing and normalization for PWD enrollment data."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


ProgressCallback = Callable[[int, int, str], None]

REPORT_COLUMNS: list[str] = [
    "Regions",
    "Province",
    "City",
    "Sector",
    "Qualification",
    "Client type",
    "Sex",
    "Age group",
    "status",
    "provider",
    "class name",
    "delivery mode",
    "TR Status",
    "Disability type",
    "Count",
]

# Aliases for metadata fields (normalized to lowercase without punctuation)
METADATA_ALIASES: dict[str, set[str]] = {
    "Regions": {"regions", "region name", "region", "reg"},
    "Province": {"province", "province name", "prov"},
    "City": {"city", "city name", "municipality", "city municipality"},
    "Sector": {"sector", "sector title", "sector name"},
    "Qualification": {"qualification", "qualification title", "qual"},
    "Client type": {"client type", "client", "clienttype"},
    "Age group": {"age group", "age", "agegroup"},
    "status": {"status"},
    "provider": {"provider", "provider name", "institution", "school"},
    "class name": {"class name", "class", "classname"},
    "delivery mode": {
        "delivery mode",
        "delivery mode title",
        "delivery mode tittle",
        "deliverymode",
    },
    "TR Status": {"tr status", "trstatus", "tr_status"},
}


def _normalize_header(text: Any) -> str:
    """Normalize a header string for alias lookup."""
    if text is None:
        return ""
    cleaned = (
        str(text)
        .lower()
        .replace("_", " ")
        .replace("-", " ")
        .replace(".", " ")
        .strip()
    )
    return " ".join(cleaned.split())


def _find_field_value(
    row_meta: dict[str, Any], target_field: str, default: str = "N/A"
) -> str:
    """Find the field value in row_meta matching aliases, or return default."""
    aliases = METADATA_ALIASES.get(target_field, set())
    for col_hdr, val in row_meta.items():
        norm = _normalize_header(col_hdr)
        if norm in aliases or norm == target_field.lower():
            if val is not None:
                str_val = str(val).strip()
                if str_val and str_val.lower() not in ("nan", "none", "null"):
                    return str_val
    return default


def format_pwd_report(
    input_file: str | Path,
    output_target: str | Path,
    sheet_name: str | None = None,
    progress_callback: ProgressCallback | None = None,
) -> Path:
    """Transform PWD enrollment data into the standardized 'Report to mimic' format.

    Args:
        input_file: Path to the input Excel workbook (.xlsx).
        output_target: Output folder or explicit .xlsx file path.
        sheet_name: Worksheet name to read (defaults to 'Sheet2' if present, else active/first).
        progress_callback: Optional progress reporter (current, total, message).

    Returns:
        Path to the saved report workbook.
    """
    input_path = Path(input_file).resolve()
    if not input_path.is_file():
        raise FileNotFoundError(f"Input file not found: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx files are supported.")

    output_dest = Path(output_target).resolve()
    if output_dest.suffix.lower() == ".xlsx":
        output_file = output_dest
        output_file.parent.mkdir(parents=True, exist_ok=True)
    else:
        output_dest.mkdir(parents=True, exist_ok=True)
        output_file = output_dest / f"Report_{input_path.stem}.xlsx"

    if progress_callback:
        progress_callback(5, 100, "Opening input workbook...")

    wb_read = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    try:
        available_sheets = wb_read.sheetnames
        chosen_sheet = None
        if sheet_name and sheet_name.strip():
            candidate = sheet_name.strip()
            if candidate in available_sheets:
                chosen_sheet = candidate
            else:
                raise ValueError(
                    f"Worksheet '{sheet_name}' not found. Available sheets: {available_sheets}"
                )
        elif "Sheet2" in available_sheets:
            chosen_sheet = "Sheet2"
        else:
            chosen_sheet = available_sheets[0]

        ws = wb_read[chosen_sheet]
        all_rows = list(ws.iter_rows(values_only=True))
    finally:
        wb_read.close()

    if len(all_rows) < 3:
        raise ValueError(
            f"The worksheet '{chosen_sheet}' must contain at least 2 header rows and at least 1 data row."
        )

    if progress_callback:
        progress_callback(20, 100, "Analyzing headers and disability columns...")

    row1 = all_rows[0]
    row2 = all_rows[1]
    data_rows = all_rows[2:]
    total_data_rows = len(data_rows)

    current_disability: str | None = None
    metadata_map: dict[int, str] = {}
    disability_cols: list[tuple[int, str, str]] = []

    for idx in range(max(len(row1), len(row2))):
        val1 = row1[idx] if idx < len(row1) else None
        val2 = row2[idx] if idx < len(row2) else None

        if val1 is not None and str(val1).strip():
            current_disability = str(val1).strip()

        sub_hdr = str(val2).strip() if val2 is not None else ""
        sub_hdr_lower = sub_hdr.lower()

        if sub_hdr_lower in ("male", "female"):
            dis_name = current_disability if current_disability else "Unspecified Disability"
            disability_cols.append((idx, dis_name, sub_hdr.title()))
        elif sub_hdr_lower == "total":
            # Skip aggregated total columns to avoid duplicate counting
            pass
        elif sub_hdr:
            metadata_map[idx] = sub_hdr

    if not disability_cols:
        raise ValueError(
            "Could not identify disaggregated disability columns (expected row 1 disability groups "
            "with row 2 'Male'/'Female' sub-headers)."
        )

    if progress_callback:
        progress_callback(35, 100, f"Unpivoting {total_data_rows} rows...")

    transformed_records: list[list[Any]] = []

    for row_idx, row in enumerate(data_rows):
        # Extract metadata dictionary for current row
        row_len = len(row)
        row_meta = {
            metadata_map[col_idx]: (row[col_idx] if col_idx < row_len else None)
            for col_idx in metadata_map
        }

        # Cache extracted metadata fields for this row
        row_fields = {
            "Regions": _find_field_value(row_meta, "Regions"),
            "Province": _find_field_value(row_meta, "Province"),
            "City": _find_field_value(row_meta, "City"),
            "Sector": _find_field_value(row_meta, "Sector"),
            "Qualification": _find_field_value(row_meta, "Qualification"),
            "Client type": _find_field_value(row_meta, "Client type", default="PWD"),
            "Age group": _find_field_value(row_meta, "Age group"),
            "status": _find_field_value(row_meta, "status"),
            "provider": _find_field_value(row_meta, "provider"),
            "class name": _find_field_value(row_meta, "class name"),
            "delivery mode": _find_field_value(row_meta, "delivery mode"),
            "TR Status": _find_field_value(row_meta, "TR Status"),
        }

        for col_idx, dis_name, sex in disability_cols:
            if col_idx >= row_len:
                continue
            val = row[col_idx]
            if val is None or val == "":
                continue

            try:
                count_val = int(val)
            except (ValueError, TypeError):
                continue

            if count_val <= 0:
                continue

            record = [
                row_fields["Regions"],
                row_fields["Province"],
                row_fields["City"],
                row_fields["Sector"],
                row_fields["Qualification"],
                row_fields["Client type"],
                sex,
                row_fields["Age group"],
                row_fields["status"],
                row_fields["provider"],
                row_fields["class name"],
                row_fields["delivery mode"],
                row_fields["TR Status"],
                dis_name,
                count_val,
            ]
            transformed_records.append(record)

        if progress_callback and (row_idx % 500 == 0 or row_idx == total_data_rows - 1):
            pct = 35 + int((row_idx / total_data_rows) * 45)
            progress_callback(
                pct,
                100,
                f"Processed {row_idx + 1} of {total_data_rows} rows ({len(transformed_records)} records)...",
            )

    if progress_callback:
        progress_callback(85, 100, f"Writing {len(transformed_records)} records to Excel...")

    df_out = pd.DataFrame(transformed_records, columns=REPORT_COLUMNS)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="Sheet1", index=False)
        ws_out = writer.sheets["Sheet1"]

        # Adjust column widths for clean readability
        for col_cells in ws_out.columns:
            first_cell = col_cells[0]
            col_letter = openpyxl.utils.get_column_letter(first_cell.column)
            # Find max length of header and top 50 sample values to set width efficiently
            max_len = len(str(first_cell.value or ""))
            for cell in col_cells[1:51]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws_out.column_dimensions[col_letter].width = max(max_len + 3, 12)

    if progress_callback:
        progress_callback(100, 100, f"Saved {len(transformed_records)} records to {output_file.name}")

    return output_file
