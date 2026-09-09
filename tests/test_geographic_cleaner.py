from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from Scripts.geographic_cleaner import clean_geographic_names


class GeographicCleanerTests(unittest.TestCase):
    def setUp(self):
        self.root = Path.cwd() / "tests" / "_runtime" / uuid.uuid4().hex
        self.root.mkdir(parents=True)
        self.metadata = self.root / "Metadata"
        self.output = self.root / "Output"
        self.metadata.mkdir()

        pd.DataFrame(
            [{"id": 4, "name": "REGION IV-A (CALABARZON)"}]
        ).to_excel(self.metadata / "regions.xlsx", index=False)
        pd.DataFrame(
            [{"id": 36, "name": "LAGUNA", "region_id": 4}]
        ).to_excel(self.metadata / "provinces.xlsx", index=False)
        pd.DataFrame(
            [
                {"id": 100, "name": "CITY OF SANTA ROSA", "province_id": 36},
                {"id": 101, "name": "SAN PEDRO", "province_id": 36},
            ]
        ).to_excel(self.metadata / "municipalities.xlsx", index=False)

    def tearDown(self):
        shutil.rmtree(self.root)

    def test_corrects_plural_headers_and_preserves_other_cells(self):
        input_file = self.root / "incoming.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Import"
        worksheet.append(["Office submission"])
        worksheet.append([])
        worksheet.append(["Regions", "Province", "Municipalities", "Total"])
        worksheet.append(["Region IV-A", "Province of Laguna", "Sta. Rosa", "=1+1"])
        workbook.save(input_file)

        result = clean_geographic_names(input_file, self.metadata, self.output)

        corrected = load_workbook(result.output_file, data_only=False)["Import"]
        self.assertEqual(corrected["A4"].value, "REGION IV-A (CALABARZON)")
        self.assertEqual(corrected["B4"].value, "LAGUNA")
        self.assertEqual(corrected["C4"].value, "CITY OF SANTA ROSA")
        self.assertEqual(corrected["D4"].value, "=1+1")
        self.assertEqual(result.cells_changed, 3)
        self.assertEqual(result.review_count, 0)
        self.assertTrue(result.report_file.is_file())

    def test_ambiguous_exact_match_is_not_changed(self):
        pd.DataFrame(
            [
                {"id": 100, "name": "SAN JOSE", "province_id": 36},
                {"id": 101, "name": "SAN JOSE", "province_id": 36},
            ]
        ).to_excel(self.metadata / "municipalities.xlsx", index=False)

        input_file = self.root / "ambiguous.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Municipality"])
        worksheet.append(["San Jose"])
        workbook.save(input_file)

        result = clean_geographic_names(input_file, self.metadata, self.output)

        corrected = load_workbook(result.output_file).active
        self.assertEqual(corrected["A2"].value, "San Jose")
        self.assertEqual(result.cells_changed, 0)
        self.assertEqual(result.review_count, 1)

    def test_requires_all_three_metadata_files(self):
        (self.metadata / "regions.xlsx").unlink()
        input_file = self.root / "incoming.xlsx"
        Workbook().save(input_file)

        with self.assertRaisesRegex(FileNotFoundError, "regions.xlsx"):
            clean_geographic_names(input_file, self.metadata, self.output)

    def test_rejects_macro_enabled_workbooks(self):
        input_file = self.root / "incoming.xlsm"
        input_file.touch()

        with self.assertRaisesRegex(ValueError, "macro-enabled"):
            clean_geographic_names(input_file, self.metadata, self.output)

    def test_does_not_fall_back_outside_matched_parent(self):
        pd.DataFrame(
            [{"id": 999, "name": "SAN PEDRO", "province_id": 999}]
        ).to_excel(self.metadata / "municipalities.xlsx", index=False)

        input_file = self.root / "wrong-parent.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Region", "Province", "Municipality"])
        worksheet.append(["REGION IV-A", "LAGUNA", "SAN PEDRO"])
        workbook.save(input_file)

        result = clean_geographic_names(input_file, self.metadata, self.output)

        corrected = load_workbook(result.output_file).active
        self.assertEqual(corrected["C2"].value, "SAN PEDRO")
        self.assertEqual(result.review_count, 1)


if __name__ == "__main__":
    unittest.main()
